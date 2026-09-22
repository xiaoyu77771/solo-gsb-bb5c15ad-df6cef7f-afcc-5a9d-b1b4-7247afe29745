"""公式解析与依赖图的核心逻辑测试。运行：python -m pytest tests/ -v"""
import io

from openpyxl import Workbook, load_workbook

from app.formula_parser import parse_formula
from app.graph import WorkbookGraph


# ---------------------------------------------------------------- 解析器
def refs(formula):
    return parse_formula(formula)


def test_plain_cell():
    r = refs("=A1+B2")
    assert [(x.col1, x.row1) for x in r] == [(1, 1), (2, 2)]
    assert all(x.sheet is None and x.kind == "cell" for x in r)


def test_cross_sheet_and_absolute():
    r = refs("=Sheet1!$A$1*2")
    assert len(r) == 1 and r[0].sheet == "Sheet1"
    assert (r[0].col1, r[0].row1) == (1, 1)


def test_quoted_sheet_name():
    r = refs("='月报 2026'!B2+1")
    assert r[0].sheet == "月报 2026"
    assert (r[0].col1, r[0].row1) == (2, 2)


def test_unquoted_cjk_sheet():
    r = refs("=参数!B2*2")
    assert r[0].sheet == "参数"


def test_rectangular_range():
    r = refs("=SUM(计算!B2:B5)")
    assert len(r) == 1
    assert r[0].kind == "range" and r[0].sheet == "计算"
    assert (r[0].col1, r[0].row1, r[0].col2, r[0].row2) == (2, 2, 2, 5)


def test_local_range_and_multiple_refs():
    r = refs("=SUM(A1:B2)+SUM(C1:C3)+D1")
    kinds = [x.kind for x in r]
    assert kinds == ["range", "range", "cell"]


def test_column_range():
    r = refs("=SUM(参数!B:B)")
    assert r[0].kind == "cols"
    assert (r[0].col1, r[0].col2) == (2, 2)
    assert r[0].row1 is None


def test_row_range():
    r = refs("=SUM(Sheet1!2:5)")
    assert r[0].kind == "rows"
    assert (r[0].row1, r[0].row2) == (2, 5)


def test_function_names_not_cells():
    # LOG10/SUM 等函数名不能被误判为单元格
    assert refs("=LOG10(A1)+SUM(B1:B2)") != []
    assert all(x.kind != "cell" or (x.col1, x.row1) != (12, 10) for x in refs("=LOG10(A1)"))
    assert len(refs("=LOG10(A1)")) == 1


def test_strings_and_errors_ignored():
    assert refs('=IF(A1>0,"B2","C3")') == refs("=IF(A1>0,1,2)")
    assert len(refs('="A1"&"B2"')) == 0
    assert refs("=#REF!+A1")[-1].row1 == 1


def test_escaped_quote_in_sheet_name():
    r = refs("='O''Brien'!A1")
    assert r[0].sheet == "O'Brien"


def test_scientific_notation_not_cell():
    # 1E2 / 1.5e-3 中的 E2、E-3 不能被当成单元格
    assert refs("=1E2+A1")[-1].row1 == 1
    assert len(refs("=1E2")) == 0
    assert len(refs("=1.5E10")) == 0
    # 紧邻单元格的数字运算仍能正确识别
    assert [(x.col1, x.row1) for x in refs("=2E3+E3")] == [(5, 3)]


def test_mixed_case_and_dollar():
    r = refs("=aa$12 + $x$9")
    assert [(x.col1, x.row1) for x in r] == [(27, 12), (24, 9)]


def test_no_ref_formula():
    assert refs("=1+2*3") == []
    assert refs('="hello"') == []


# ---------------------------------------------------------------- 依赖图
def build_graph():
    """构造：参数!B1 -> 计算!B1 -> 计算!B2 -> 汇总!B1；循环!A1 <-> 循环!B1；循环!C1 自环"""
    wb = Workbook()
    p = wb.active
    p.title = "参数"
    p["B1"] = 10
    c = wb.create_sheet("计算")
    c["B1"] = "=参数!B1*2"
    c["B2"] = "=B1+SUM(参数!B1:B2)"
    s = wb.create_sheet("汇总")
    s["B1"] = "=计算!B2"
    s["B2"] = "=SUM(计算!B1:B2)"
    loop = wb.create_sheet("循环")
    loop["A1"] = "=B1+1"
    loop["B1"] = "=A1+1"
    loop["C1"] = "=C1*2"
    loop["D1"] = "=A1+5"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    wbf = load_workbook(buf, data_only=False)
    buf.seek(0)
    wbv = load_workbook(buf, data_only=True)
    return WorkbookGraph.from_workbooks(wbf, wbv)


def test_graph_edges():
    g = build_graph()
    assert "参数!B1" in g.dependents
    assert "计算!B1" in g.dependents["参数!B1"]
    assert "计算!B2" in g.dependents["计算!B1"]
    assert "汇总!B1" in g.dependents["计算!B2"]
    # 区域 SUM(计算!B1:B2) 展开为两条边
    assert "汇总!B2" in g.dependents["计算!B1"]
    assert "汇总!B2" in g.dependents["计算!B2"]


def test_cycle_detection():
    g = build_graph()
    flat = set()
    for cyc in g.cycles:
        flat.update(cyc)
    assert {"循环!A1", "循环!B1"} <= flat      # 双向环
    assert "循环!C1" in flat                    # 自环
    assert "计算!B1" not in flat                # 正常节点不入环
    assert "循环!D1" not in flat                # 环的下游本身不在环里


def test_impact_downstream():
    g = build_graph()
    res = g.impact("参数!B1", direction="downstream")
    cells = {i["cell"] for i in res["impacted"]}
    assert {"计算!B1", "计算!B2", "汇总!B1", "汇总!B2"} <= cells
    # 路径应从根出发
    b2 = next(i for i in res["impacted"] if i["cell"] == "汇总!B1")
    assert b2["path"][0] == "参数!B1" and b2["path"][-1] == "汇总!B1"
    assert b2["depth"] == 2  # B2 经 SUM(参数!B1:B2) 也直接引用 参数!B1，最短路径深度为 2


def test_impact_upstream():
    g = build_graph()
    res = g.impact("汇总!B1", direction="upstream")
    cells = {i["cell"] for i in res["impacted"]}
    assert {"计算!B2", "计算!B1", "参数!B1"} <= cells


def test_impact_direct():
    g = build_graph()
    res = g.impact("计算!B1", direction="downstream")
    direct = {d["cell"] for d in res["direct"]}
    assert direct == {"计算!B2", "汇总!B2"}


def test_export_structure():
    g = build_graph()
    data = g.export("t.xlsx")
    assert data["summary"]["cycle_count"] >= 2
    assert any(e["from"] == "计算!B2" and e["to"] == "汇总!B1" for e in data["edges"])
    assert "计算!B1" in data["formulas"]
