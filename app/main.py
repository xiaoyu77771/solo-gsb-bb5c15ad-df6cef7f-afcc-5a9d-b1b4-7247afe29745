"""FastAPI 入口：上传 XLSX、依赖图查询、影响面分析、JSON 导出。"""
from __future__ import annotations

import io
import time
import uuid
from collections import OrderedDict
from pathlib import Path

from fastapi import FastAPI, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, JSONResponse
from openpyxl import load_workbook

from .graph import (
    AUDIT_LEVELS,
    DEFAULT_FANOUT_THRESHOLD,
    WorkbookGraph,
    display_id,
)

MAX_UPLOAD_BYTES = 25 * 1024 * 1024  # 25 MB
MAX_WORKBOOKS = 5  # 内存中最多保留的工作簿数量，超出后淘汰最旧的

STATIC_DIR = Path(__file__).parent / "static"
SAMPLES_DIR = Path(__file__).resolve().parent.parent / "samples"

app = FastAPI(title="电子表格公式影响分析", version="1.1.0")

# wid -> {"graph": WorkbookGraph, "filename": str, "ts": float}
_store: OrderedDict[str, dict] = OrderedDict()


def _get(wid: str) -> dict:
    entry = _store.get(wid)
    if entry is None:
        raise HTTPException(status_code=404, detail="工作簿不存在或已被淘汰，请重新上传")
    _store.move_to_end(wid)
    return entry


@app.get("/")
def index():
    return FileResponse(STATIC_DIR / "index.html")


@app.post("/api/upload")
async def upload(file: UploadFile = File(...)):
    filename = file.filename or "workbook.xlsx"
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx / .xlsm 文件")
    data = await file.read()
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="文件超过 25MB 限制")
    try:
        wb_formula = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
        wb_values = load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"无法解析文件: {exc}") from exc

    graph = WorkbookGraph.from_workbooks(wb_formula, wb_values)
    wid = uuid.uuid4().hex[:12]
    _store[wid] = {"graph": graph, "filename": filename, "ts": time.time()}
    while len(_store) > MAX_WORKBOOKS:
        _store.popitem(last=False)

    return {"workbook_id": wid, "filename": filename, **graph.summary()}


@app.get("/api/workbooks/{wid}")
def workbook_summary(wid: str):
    entry = _get(wid)
    return {
        "workbook_id": wid,
        "filename": entry["filename"],
        **entry["graph"].summary(),
    }


@app.get("/api/workbooks/{wid}/cells")
def workbook_cells(
    wid: str,
    sheet: str | None = Query(default=None),
    q: str = Query(default=""),
    formulas_only: bool = Query(default=True),
    limit: int = Query(default=1000, ge=1, le=10000),
    offset: int = Query(default=0, ge=0),
):
    graph = _get(wid)["graph"]
    return graph.cell_list(sheet=sheet, query=q, formulas_only=formulas_only,
                           limit=limit, offset=offset)


@app.get("/api/workbooks/{wid}/impact")
def workbook_impact(
    wid: str,
    cell: str = Query(..., description="形如 '计算!B1' 的单元格 id"),
    direction: str = Query(default="downstream", pattern="^(downstream|upstream)$"),
    max_depth: int = Query(default=50, ge=1, le=200),
):
    graph = _get(wid)["graph"]
    # 兼容带引号的展示形式：'月报 2026'!B1 -> 月报 2026!B1
    normalized = cell
    if normalized.startswith("'") and "'!" in normalized:
        normalized = normalized[1:].replace("'!", "!", 1)
    if "!" not in normalized:
        raise HTTPException(status_code=400, detail="cell 参数格式应为 工作表!坐标")
    sheet, coord = normalized.rsplit("!", 1)
    coord = coord.replace("$", "").upper()
    if sheet not in {s.name for s in graph.sheets}:
        raise HTTPException(status_code=404, detail=f"工作表不存在: {sheet}")
    node = f"{sheet}!{coord}"
    result = graph.impact(node, direction=direction, max_depth=max_depth)
    result["root_display"] = display_id(sheet, coord)
    result["root_formula"] = graph.formulas.get(node)
    result["root_value"] = graph.values.get(node)
    result["root_in_cycle"] = node in graph.cycle_nodes
    return result


@app.get("/api/workbooks/{wid}/audit")
def workbook_audit(
    wid: str,
    level: list[str] | None = Query(
        default=None,
        description="问题级别筛选，可重复传参：error / warning / info"),
    sheet: str | None = Query(default=None, description="按工作表筛选"),
    fanout_threshold: str | None = Query(
        default=None,
        description=f"扇出告警阈值，1-10000 的整数，默认 {DEFAULT_FANOUT_THRESHOLD}"),
):
    graph = _get(wid)["graph"]

    levels = AUDIT_LEVELS if not level else tuple(dict.fromkeys(level))
    bad = [lv for lv in levels if lv not in AUDIT_LEVELS]
    if bad:
        raise HTTPException(
            status_code=400,
            detail=f"level 取值非法: {bad[0]}，仅支持 {list(AUDIT_LEVELS)}")
    if sheet is not None and sheet not in {s.name for s in graph.sheets}:
        raise HTTPException(status_code=404, detail=f"工作表不存在: {sheet}")

    threshold = DEFAULT_FANOUT_THRESHOLD
    if fanout_threshold is not None:
        try:
            threshold = int(fanout_threshold)
        except (TypeError, ValueError):
            raise HTTPException(
                status_code=400,
                detail="fanout_threshold 必须是 1-10000 的整数") from None
        if not 1 <= threshold <= 10000:
            raise HTTPException(
                status_code=400,
                detail="fanout_threshold 必须是 1-10000 的整数")

    return {
        "workbook_id": wid,
        **graph.audit(levels=levels, sheet=sheet, fanout_threshold=threshold),
    }


@app.get("/api/workbooks/{wid}/audit/export")
def workbook_audit_export(wid: str):
    entry = _get(wid)
    payload = {
        "workbook_id": wid,
        "filename": entry["filename"],
        "audit": entry["graph"].audit(evidence_limit=None),
    }
    return JSONResponse(
        payload,
        headers={
            "Content-Disposition": (
                f"attachment; filename=dependency_audit_{wid}.json; "
                f"filename*=UTF-8''dependency_audit_{wid}.json"
            )
        },
    )


@app.get("/api/workbooks/{wid}/export")
def workbook_export(wid: str):
    entry = _get(wid)
    payload = entry["graph"].export(entry["filename"])
    return JSONResponse(
        payload,
        headers={
            "Content-Disposition": (
                f"attachment; filename=impact_analysis_{wid}.json; "
                f"filename*=UTF-8''impact_analysis_{wid}.json"
            )
        },
    )


@app.get("/samples/{name}")
def sample_file(name: str):
    """提供 samples/ 目录下的示例文件下载。"""
    path = (SAMPLES_DIR / name).resolve()
    if not str(path).startswith(str(SAMPLES_DIR.resolve())) or not path.is_file():
        raise HTTPException(status_code=404, detail="示例文件不存在")
    return FileResponse(path, filename=name)


@app.get("/api/health")
def health():
    return {"status": "ok", "workbooks": len(_store)}
